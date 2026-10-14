from io import BytesIO
from zipfile import ZipFile, BadZipFile
import sqlite3
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import ParseError
from ..models.database import get_db
from .admision_service import datos_paciente, validar_admision, insertar_admision, evento

COLUMNAS=['documento','nombre','nacimiento','telefono','correo','fecha','especialidad','motivo','profesional_id']

def leer_excel(stream):
    data=stream.read(2*1024*1024+1)
    if len(data)>2*1024*1024:
        raise ValueError('El archivo supera 2 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries=archive.infolist()
            if len(entries)>1000 or sum(i.file_size for i in entries)>20*1024*1024:
                raise ValueError('El Excel supera el límite de contenido descomprimido.')
            if any('vbaproject' in i.filename.lower() for i in entries):
                raise ValueError('No se admiten macros.')
        wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            if len(wb.worksheets)!=1:
                raise ValueError('Usa una sola hoja y los encabezados de la plantilla.')
            ws=wb.worksheets[0]
            if ws.max_row is not None and ws.max_row>1001:
                raise ValueError('Máximo 1000 filas de pacientes por archivo.')
            if ws.max_column is not None and ws.max_column>len(COLUMNAS):
                raise ValueError('El archivo contiene columnas extra.')
            headers=[str(v or '').strip().lower() for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
            if headers!=COLUMNAS:
                raise ValueError('Encabezados incorrectos. Descarga la plantilla desde Admisión.')
            rows=[]
            for numero,cells in enumerate(ws.iter_rows(min_row=2,max_row=1002,max_col=len(COLUMNAS)),start=2):
                if not any(c.value is not None for c in cells):
                    continue
                if numero>1001:
                    raise ValueError('Máximo 1000 filas por archivo.')
                if any(c.data_type in ('f','e') for c in cells):
                    raise ValueError(f'Fila {numero}: no se admiten fórmulas ni errores de Excel.')
                rows.append((numero,dict(zip(COLUMNAS,[c.value for c in cells]))))
            if not rows:
                raise ValueError('El archivo no contiene pacientes.')
            return rows
        finally:
            wb.close()
    except (BadZipFile,InvalidFileException,DefusedXmlException,ParseError,KeyError,OSError) as e:
        raise ValueError('El archivo no es un Excel válido o seguro.') from e

def importar_filas(rows,actor):
    db=get_db()
    errores=[]
    with db:
        db.execute('BEGIN IMMEDIATE')
        for numero,row in rows:
            db.execute('SAVEPOINT fila_excel')
            try:
                valores=datos_paciente(row)
                paciente=db.execute('SELECT * FROM pacientes WHERE documento=?',(valores[0],)).fetchone()
                if paciente:
                    if not paciente['activo']:
                        raise ValueError('Paciente desactivado; no se reactiva por importación.')
                    if paciente['nombre'].casefold()!=valores[1].casefold() or paciente['nacimiento']!=valores[2]:
                        raise ValueError('El documento existe con nombre o nacimiento diferente. Corrige el archivo o edita el paciente.')
                else:
                    c=db.execute('INSERT INTO pacientes(documento,nombre,nacimiento,telefono,correo) VALUES(?,?,?,?,?)',valores)
                    paciente=db.execute('SELECT * FROM pacientes WHERE id=?',(c.lastrowid,)).fetchone()
                    evento(db,actor,'paciente_importado',paciente['id'],'paciente')
                insertar_admision(db,paciente['id'],validar_admision(row,paciente),actor)
            except ValueError as e:
                db.execute('ROLLBACK TO fila_excel')
                errores.append(f'Fila {numero}: {e}')
            except sqlite3.IntegrityError:
                db.execute('ROLLBACK TO fila_excel')
                errores.append(f'Fila {numero}: documento repetido o paciente con admisión abierta.')
            finally:
                db.execute('RELEASE fila_excel')
        if errores:
            raise ValueError('No se importó ninguna fila. Corrige estos errores:\n'+'\n'.join(errores[:30]))
        evento(db,actor,'excel_importado',None)
    return len(rows)

def crear_excel(headers,rows,titulo='Listado'):
    wb=Workbook()
    ws=wb.active;ws.title=titulo
    ws.append(headers)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='124E61')
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):
                cell.data_type='s'
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(40,max(18,len(str(col[0].value))+3))
    output=BytesIO();wb.save(output);wb.close();output.seek(0)
    return output
