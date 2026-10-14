from io import BytesIO
import sqlite3
import unittest
from openpyxl import Workbook, load_workbook
import test_clinica
from app.models.database import get_db
from app.models.usuario_model import crear_usuario
from app.services.excel_service import COLUMNAS

class AdmisionTest(unittest.TestCase):
    def setUp(self):
        self.b=test_clinica.ClinicaTest();self.b.setUp()
    def tearDown(self):
        self.b.tearDown()
    def nueva(self,documento='FICTICIO-1',pid='1'):
        self.b.login('recepcion')
        return self.b.post('/admisiones/nueva',documento=documento,nombre='Paciente de prueba',nacimiento='2000-01-01',fecha='2020-01-01',especialidad='Medicina',motivo='Control',profesional_id=pid)
    def excel(self,rows):
        wb=Workbook();ws=wb.active;ws.append(COLUMNAS)
        for row in rows: ws.append(row)
        out=BytesIO();wb.save(out);wb.close();out.seek(0);return out
    def upload(self,rows):
        self.b.login('recepcion')
        return self.b.post('/admisiones/importar',archivo=(self.excel(rows),'prueba.xlsx'))
    def test_crud_y_control_de_version(self):
        self.assertEqual(self.nueva().status_code,302)
        self.assertEqual(self.b.get('/admisiones').status_code,200)
        self.assertEqual(self.b.get('/admisiones/1/editar').status_code,200)
        self.assertEqual(self.b.post('/admisiones/1/editar',version='1',fecha='2020-01-02',especialidad='Medicina',motivo='Nuevo motivo',profesional_id='1',estado='pendiente').status_code,302)
        self.assertEqual(self.b.post('/admisiones/1/editar',version='1',fecha='2020-01-03',especialidad='Medicina',motivo='Obsoleto',profesional_id='1',estado='pendiente').status_code,409)
        self.assertEqual(self.b.get('/admisiones/1/eliminar').status_code,200)
        self.assertEqual(self.b.post('/admisiones/1/eliminar',confirmar='ELIMINAR',version='2').status_code,302)
        self.assertEqual(self.b.get('/admisiones/1/editar').status_code,404)
        self.assertEqual(self.b.get('/pacientes/1').status_code,200)
    def test_atencion_asignada_guarda_historial_una_vez(self):
        self.nueva()
        self.b.login('clinicob')
        self.assertEqual(self.b.get('/admisiones/1/atender').status_code,404)
        self.b.login('clinicoa')
        self.assertEqual(self.b.get('/admisiones/1/atender').status_code,200)
        self.assertEqual(self.b.post('/admisiones/1/atender',version='1',accion='iniciar').status_code,302)
        self.assertEqual(self.b.post('/admisiones/1/atender',version='2',accion='finalizar',motivo='Control',diagnostico='Ejemplo',tratamiento='Indicaciones').status_code,302)
        self.assertEqual(self.b.post('/admisiones/1/atender',version='2',accion='finalizar',motivo='X',diagnostico='X',tratamiento='X').status_code,409)
        with self.b.app.app_context():
            db=get_db()
            self.assertEqual(db.execute('SELECT estado FROM admisiones').fetchone()[0],'atendido')
            self.assertEqual(db.execute('SELECT COUNT(*) FROM atenciones').fetchone()[0],1)
            self.assertEqual(db.execute('SELECT admision_id FROM atenciones').fetchone()[0],1)
    def test_portal_paciente_aislado(self):
        self.nueva();self.nueva('FICTICIO-2')
        with self.b.app.app_context():
            crear_usuario('paciente1','Paciente','Frase de prueba larga','paciente',paciente_id=1)
        self.b.login('paciente1')
        self.assertEqual(self.b.get('/mis-admisiones').status_code,200)
        self.assertEqual(self.b.get('/mis-admisiones/1').status_code,200)
        self.assertEqual(self.b.get('/mis-admisiones/2').status_code,404)
        for path in ['/admisiones','/pacientes/1','/usuarios','/admisiones/plantilla.xlsx']:
            self.assertEqual(self.b.get(path).status_code,403)
        self.assertEqual(self.b.post('/admisiones/nueva').status_code,403)
    def test_excel_valido_y_repetido(self):
        rows=[['EXCEL-1','Ejemplo Excel','2000-01-01','','','2020-01-01','Medicina','Control',1]]
        self.assertEqual(self.upload(rows).status_code,302)
        self.assertEqual(self.upload(rows).status_code,200)
        with self.b.app.app_context():
            db=get_db()
            self.assertEqual(db.execute('SELECT COUNT(*) FROM pacientes').fetchone()[0],1)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM admisiones').fetchone()[0],1)
        self.assertEqual(self.b.get('/admisiones/plantilla.xlsx').status_code,200)
        r=self.b.get('/admisiones/exportar.xlsx')
        wb=load_workbook(BytesIO(r.data));self.assertEqual(wb.active.cell(2,2).value,'EXCEL-1');wb.close()
    def test_excel_invalido_rechaza_todo_y_formula(self):
        valid=['EXCEL-1','Primero','2000-01-01','','','2020-01-01','Medicina','Control',1]
        bad=['EXCEL-2','Segundo','2999-01-01','','','2020-01-01','Medicina','Control',1]
        self.assertEqual(self.upload([valid,bad]).status_code,200)
        with self.b.app.app_context():
            db=get_db()
            self.assertEqual(db.execute('SELECT COUNT(*) FROM pacientes').fetchone()[0],0)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM admisiones').fetchone()[0],0)
        formula=valid.copy();formula[1]='=1+1'
        self.assertEqual(self.upload([formula]).status_code,200)
        with self.b.app.app_context():
            self.assertEqual(get_db().execute('SELECT COUNT(*) FROM pacientes').fetchone()[0],0)
    def test_edicion_pacientes_funcionarios_y_cuentas(self):
        self.nueva()
        self.assertEqual(self.b.get('/pacientes/1/editar').status_code,200)
        self.assertEqual(self.b.post('/pacientes/1/editar',documento='FICTICIO-1',nombre='Nombre editado',nacimiento='2000-01-01').status_code,302)
        self.b.login('admin')
        self.assertEqual(self.b.get('/profesionales/1/editar').status_code,200)
        self.assertEqual(self.b.post('/profesionales/1/editar',nombre='Profesional editado',especialidad='Medicina',cargo='Médico').status_code,302)
        self.assertEqual(self.b.post('/profesionales/1/eliminar',confirmar='ELIMINAR').status_code,200)
        self.assertEqual(self.b.get('/usuarios/2/editar').status_code,200)
        self.assertEqual(self.b.post('/usuarios/2/editar',username='recepcion',nombre='Recepción nueva',rol='recepcion').status_code,302)
        self.assertEqual(self.b.post('/usuarios/1/editar',username='admin',nombre='Admin',rol='paciente',paciente_id='1').status_code,200)
    def test_migracion_cuentas_antiguas_con_sesiones(self):
        with self.b.app.app_context():
            db=get_db();db.commit();db.execute('PRAGMA foreign_keys=OFF')
            db.execute('CREATE TABLE usuarios_viejos(id INTEGER PRIMARY KEY,username TEXT NOT NULL UNIQUE,nombre TEXT NOT NULL,password_hash TEXT NOT NULL,rol TEXT NOT NULL CHECK(rol IN (\'admin\',\'recepcion\',\'profesional\',\'auditor\')),profesional_id INTEGER,activo INTEGER NOT NULL DEFAULT 1)')
            db.execute('INSERT INTO usuarios_viejos SELECT id,username,nombre,password_hash,rol,profesional_id,activo FROM usuarios')
            db.execute('DROP TABLE usuarios');db.execute('ALTER TABLE usuarios_viejos RENAME TO usuarios');db.commit();db.execute('PRAGMA foreign_keys=ON')
        from app.models.database import init_db
        init_db(self.b.app)
        self.assertEqual(self.b.login('admin').status_code,302)
        with self.b.app.app_context():
            self.assertIsNone(get_db().execute('PRAGMA foreign_key_check').fetchone())

    def test_mismo_paciente_atendido_por_dos_funcionarios(self):
        self.nueva()
        self.b.login('clinicoa')
        self.b.post('/admisiones/1/atender',version='1',accion='iniciar')
        self.b.post('/admisiones/1/atender',version='2',accion='finalizar',motivo='Primero',diagnostico='Ejemplo',tratamiento='Indicaciones')
        self.b.login('recepcion')
        self.assertEqual(self.b.post('/admisiones/nueva',paciente_id='1',fecha='2020-01-01',especialidad='Enfermería',motivo='Segundo',profesional_id='2').status_code,302)
        self.b.login('clinicob')
        self.b.post('/admisiones/2/atender',version='1',accion='iniciar')
        self.b.post('/admisiones/2/atender',version='2',accion='finalizar',motivo='Segundo',diagnostico='Ejemplo',tratamiento='Indicaciones')
        with self.b.app.app_context():
            self.assertEqual([r[0] for r in get_db().execute('SELECT profesional_id FROM atenciones ORDER BY id')],[1,2])
        r=self.b.get('/pacientes/1')
        self.assertIn(b'Profesional A',r.data);self.assertIn(b'Profesional B',r.data)
